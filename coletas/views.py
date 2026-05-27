"""
Views/ViewSets for coletas app
"""
from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated, AllowAny
from django.db import transaction
from django.utils import timezone
from .models import (
    Coleta, ColetaProfissional, RespostaColeta, Notificacao,
    StatusColeta, StatusPreenchimento
)
from .serializers import (
    ColetaSerializer, ColetaListSerializer, ColetaIniciarSerializer,
    ColetaProfissionalSerializer, RespostaColetaSerializer,
    RespostaColetaCreateSerializer, NotificacaoSerializer,
    ColetaImportExcelSerializer, ColetaEstatisticasSerializer
)


class ColetaViewSet(viewsets.ModelViewSet):
    """ViewSet para Coleta"""
    queryset = Coleta.objects.select_related('empresa', 'modelo', 'iniciado_por').all()
    serializer_class = ColetaSerializer
    permission_classes = [IsAuthenticated]
    
    def get_serializer_class(self):
        if self.action == 'list':
            return ColetaListSerializer
        elif self.action == 'iniciar':
            return ColetaIniciarSerializer
        elif self.action == 'import_excel':
            return ColetaImportExcelSerializer
        return ColetaSerializer
    
    def get_queryset(self):
        user = self.request.user
        queryset = Coleta.objects.select_related('empresa', 'modelo', 'iniciado_por').all()

        # Support both DRF Request (has .query_params) and plain Django WSGIRequest
        params = getattr(self.request, 'query_params', None)
        if params is None:
            params = getattr(self.request, 'GET', {})

        # Filtrar por empresa ou profissional
        if user.tipo_usuario == 'empresa':
            from empresas.models import Empresa
            try:
                empresa = Empresa.objects.get(usuario=user)
                queryset = queryset.filter(empresa=empresa)
            except Empresa.DoesNotExist:
                return Coleta.objects.none()
        elif user.tipo_usuario == 'profissional':
            # Profissional pode ver coletas da empresa vinculada
            from empresas.models import Profissional
            try:
                profissional = Profissional.objects.get(usuario=user)
                queryset = queryset.filter(empresa=profissional.empresa)
            except Profissional.DoesNotExist:
                return Coleta.objects.none()

        # Aplicar filtros adicionais
        empresa_id = params.get('empresa_id')
        status_coleta = params.get('status')
        modelo_id = params.get('modelo_id')
        is_active = params.get('is_active')

        if empresa_id:
            queryset = queryset.filter(empresa_id=empresa_id)
        if status_coleta:
            queryset = queryset.filter(status=status_coleta)
        if modelo_id:
            queryset = queryset.filter(modelo_id=modelo_id)
        if is_active is not None:
            queryset = queryset.filter(is_active=is_active.lower() == 'true')

        return queryset
    
    @action(detail=True, methods=['post'])
    @transaction.atomic
    def iniciar(self, request, pk=None):
        """Iniciar coleta adicionando profissionais"""
        coleta = self.get_object()
        serializer = ColetaIniciarSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        
        from empresas.models import Profissional, Departamento
        
        profissionais_ids = serializer.validated_data.get('profissionais_ids', [])
        departamentos_ids = serializer.validated_data.get('departamentos_ids', [])
        todos_profissionais = serializer.validated_data.get('todos_profissionais', False)
        enviar_notificacao = serializer.validated_data.get('enviar_notificacao', True)
        
        profissionais = []
        
        if todos_profissionais:
            profissionais = list(coleta.empresa.profissionais.filter(is_active=True))
        else:
            if profissionais_ids:
                profissionais.extend(
                    Profissional.objects.filter(usuario__in=profissionais_ids, is_active=True)
                )
            if departamentos_ids:
                for dept_id in departamentos_ids:
                    profissionais.extend(
                        Profissional.objects.filter(
                            departamento_id=dept_id,
                            is_active=True
                        )
                    )
        
        # Adicionar profissionais à coleta
        adicionados = 0
        for profissional in set(profissionais):  # Remove duplicados
            coleta.adicionar_profissional(profissional)
            adicionados += 1
            
            # TODO: Enviar notificação via Celery
            if enviar_notificacao:
                pass
        
        # Atualizar totais após adicionar profissionais
        if adicionados > 0:
            coleta.atualizar_totais()
        
        # Atualizar status da coleta
        coleta.status = StatusColeta.EM_ANDAMENTO
        if not coleta.data_inicio:
            coleta.data_inicio = timezone.now()
        coleta.save()
        
        # Registrar no log de auditoria
        from auditoria.models import LogAuditoria, TipoAcao
        from empresas.models import Empresa
        
        try:
            empresa = Empresa.objects.get(usuario=request.user)
            LogAuditoria.registrar(
                usuario=request.user,
                tipo_acao=TipoAcao.COLETA_INICIADA,
                descricao=f"Coleta '{coleta.nome}' iniciada com {adicionados} profissional(is)",
                empresa=empresa,
                coleta=coleta,
                dados_adicionais={'quantidade_profissionais': adicionados},
                request=request
            )
        except Empresa.DoesNotExist:
            pass
        
        return Response({
            'message': f'{adicionados} profissional(is) adicionado(s) à coleta.',
            'coleta': ColetaSerializer(coleta).data
        })
    
    @action(detail=True, methods=['post'])
    @transaction.atomic
    def adicionar_profissionais(self, request, pk=None):
        """Adicionar profissionais a uma coleta existente"""
        coleta = self.get_object()
        serializer = ColetaIniciarSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        
        from empresas.models import Profissional, Departamento
        
        profissionais_ids = serializer.validated_data.get('profissionais_ids', [])
        departamentos_ids = serializer.validated_data.get('departamentos_ids', [])
        todos_profissionais = serializer.validated_data.get('todos_profissionais', False)
        enviar_notificacao = serializer.validated_data.get('enviar_notificacao', True)
        
        profissionais = []
        
        if todos_profissionais:
            # Buscar profissionais que ainda não estão na coleta
            profissionais_existentes = coleta.coleta_profissionais.values_list('profissional_id', flat=True)
            profissionais = list(
                coleta.empresa.profissionais.filter(is_active=True).exclude(
                    usuario__in=profissionais_existentes
                )
            )
        else:
            if profissionais_ids:
                profissionais_existentes = coleta.coleta_profissionais.values_list('profissional_id', flat=True)
                profissionais.extend(
                    Profissional.objects.filter(
                        usuario__in=profissionais_ids, 
                        is_active=True
                    ).exclude(usuario__in=profissionais_existentes)
                )
            if departamentos_ids:
                profissionais_existentes = coleta.coleta_profissionais.values_list('profissional_id', flat=True)
                for dept_id in departamentos_ids:
                    profissionais.extend(
                        Profissional.objects.filter(
                            departamento_id=dept_id,
                            is_active=True
                        ).exclude(usuario__in=profissionais_existentes)
                    )
        
        # Adicionar profissionais à coleta
        adicionados = 0
        for profissional in set(profissionais):  # Remove duplicados
            coleta.adicionar_profissional(profissional)
            adicionados += 1
            
            # TODO: Enviar notificação via Celery
            if enviar_notificacao:
                pass
        
        # Atualizar totais após adicionar profissionais
        if adicionados > 0:
            coleta.atualizar_totais()
        
        # Registrar no log de auditoria
        if adicionados > 0:
            from auditoria.models import LogAuditoria, TipoAcao
            from empresas.models import Empresa
            
            try:
                empresa = Empresa.objects.get(usuario=request.user)
                LogAuditoria.registrar(
                    usuario=request.user,
                    tipo_acao=TipoAcao.PARTICIPANTE_ADICIONADO,
                    descricao=f"{adicionados} profissional(is) adicionado(s) à coleta '{coleta.nome}'",
                    empresa=empresa,
                    coleta=coleta,
                    dados_adicionais={'quantidade': adicionados},
                    request=request
                )
            except Empresa.DoesNotExist:
                pass
        
        return Response({
            'message': f'{adicionados} profissional(is) adicionado(s) à coleta.',
            'coleta': ColetaSerializer(coleta).data
        })
    
    @action(detail=True, methods=['get'])
    def exportar_excel(self, request, pk=None):
        """Exportar dados da coleta para Excel"""
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
            from django.http import HttpResponse
            from datetime import datetime
            
            coleta = self.get_object()
            
            # Criar workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Dados da Coleta"
            
            # Estilo do cabeçalho
            header_fill = PatternFill(start_color="0066CC", end_color="0066CC", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            header_alignment = Alignment(horizontal="center", vertical="center")
            
            # Informações da Coleta
            ws['A1'] = 'DADOS DA COLETA'
            ws['A1'].font = Font(bold=True, size=14)
            ws.merge_cells('A1:E1')
            
            ws['A3'] = 'Nome da Coleta:'
            ws['B3'] = coleta.nome
            ws['A4'] = 'Empresa:'
            ws['B4'] = coleta.empresa.razao_social if coleta.empresa else 'N/A'
            ws['A5'] = 'Status:'
            ws['B5'] = coleta.get_status_display()
            ws['A6'] = 'Data de Início:'
            ws['B6'] = coleta.data_inicio.strftime('%d/%m/%Y %H:%M') if coleta.data_inicio else 'Não iniciada'
            ws['A7'] = 'Total de Profissionais:'
            ws['B7'] = coleta.total_profissionais
            ws['A8'] = 'Total Concluídos:'
            ws['B8'] = coleta.total_concluidos
            ws['A9'] = 'Percentual de Conclusão:'
            ws['B9'] = f"{coleta.calcular_percentual()}%"
            ws['A10'] = 'Gerado em:'
            ws['B10'] = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
            
            # Cabeçalho da tabela de profissionais
            row = 12
            headers = ['Profissional', 'CPF', 'Status', 'Data de Conclusão']
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment
            
            # Dados dos profissionais
            row += 1
            for cp in coleta.coleta_profissionais.select_related('profissional').all():
                ws.cell(row=row, column=1, value=cp.profissional.nome if cp.profissional else 'N/A')
                ws.cell(row=row, column=2, value=cp.profissional.cpf if cp.profissional else 'N/A')
                ws.cell(row=row, column=3, value=cp.get_status_display())
                ws.cell(row=row, column=4, value=cp.data_conclusao.strftime('%d/%m/%Y %H:%M') if cp.data_conclusao else '-')
                row += 1
            
            # Ajustar largura das colunas
            ws.column_dimensions['A'].width = 30
            ws.column_dimensions['B'].width = 20
            ws.column_dimensions['C'].width = 20
            ws.column_dimensions['D'].width = 20
            
            # Preparar resposta HTTP
            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            filename = f"coleta_{coleta.nome.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            
            wb.save(response)
            return response
            
        except Exception as e:
            import traceback, logging
            logger = logging.getLogger(__name__)
            logger.error(f"[ERROR] Erro ao exportar Excel: {str(e)}")
            logger.error(traceback.format_exc())
            return Response(
                {'error': f'Erro ao exportar: {str(e)}'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )
    
    @action(detail=True, methods=['post'])
    def cancelar(self, request, pk=None):
        """Cancelar coleta"""
        coleta = self.get_object()
        coleta.cancelar()
        
        # Registrar no log de auditoria
        from auditoria.models import LogAuditoria, TipoAcao
        from empresas.models import Empresa
        
        try:
            empresa = Empresa.objects.get(usuario=request.user)
            LogAuditoria.registrar(
                usuario=request.user,
                tipo_acao=TipoAcao.COLETA_CANCELADA,
                descricao=f"Coleta '{coleta.nome}' cancelada",
                empresa=empresa,
                coleta=coleta,
                request=request
            )
        except Empresa.DoesNotExist:
            pass
        
        return Response({'message': 'Coleta cancelada com sucesso.'})
    
    @action(detail=True, methods=['post'])
    def finalizar(self, request, pk=None):
        """Finalizar coleta"""
        coleta = self.get_object()
        coleta.finalizar()
        
        # Registrar no log de auditoria
        from auditoria.models import LogAuditoria, TipoAcao
        from empresas.models import Empresa
        
        try:
            empresa = Empresa.objects.get(usuario=request.user)
            LogAuditoria.registrar(
                usuario=request.user,
                tipo_acao=TipoAcao.COLETA_CONCLUIDA,
                descricao=f"Coleta '{coleta.nome}' concluída (forçada)",
                empresa=empresa,
                coleta=coleta,
                request=request
            )
        except Empresa.DoesNotExist:
            pass
        
        return Response({'message': 'Coleta finalizada com sucesso.'})
    
    @action(detail=True, methods=['post'])
    def reabrir(self, request, pk=None):
        """Reabrir coleta"""
        coleta = self.get_object()
        coleta.reabrir()
        
        # Registrar no log de auditoria
        from auditoria.models import LogAuditoria, TipoAcao
        from empresas.models import Empresa
        
        try:
            empresa = Empresa.objects.get(usuario=request.user)
            LogAuditoria.registrar(
                usuario=request.user,
                tipo_acao=TipoAcao.COLETA_REABERTA,
                descricao=f"Coleta '{coleta.nome}' reaberta",
                empresa=empresa,
                coleta=coleta,
                request=request
            )
        except Empresa.DoesNotExist:
            pass
        
        return Response({'message': 'Coleta reaberta com sucesso.'})
    
    @action(detail=True, methods=['get'])
    def estatisticas(self, request, pk=None):
        """Estatísticas da coleta"""
        coleta = self.get_object()
        
        from django.db.models import Count, Avg, F
        
        # Recalcular totais para garantir dados atualizados
        total_profissionais = coleta.coleta_profissionais.count()
        total_concluidos = coleta.coleta_profissionais.filter(
            status=StatusPreenchimento.CONCLUIDO
        ).count()
        total_pendentes = coleta.coleta_profissionais.filter(
            status=StatusPreenchimento.PENDENTE
        ).count()
        total_em_progresso = coleta.coleta_profissionais.filter(
            status=StatusPreenchimento.EM_PROGRESSO
        ).count()
        
        # Calcular percentual
        if total_profissionais > 0:
            percentual_conclusao = round((total_concluidos / total_profissionais) * 100, 2)
        else:
            percentual_conclusao = 0.0
        
        profissionais_por_status = ColetaProfissional.objects.filter(
            coleta=coleta
        ).values('status').annotate(total=Count('id'))
        
        tempo_medio = ColetaProfissional.objects.filter(
            coleta=coleta,
            status=StatusPreenchimento.CONCLUIDO
        ).aggregate(
            tempo_medio=Avg(F('data_conclusao') - F('data_inicio'))
        )
        
        data = {
            'total_profissionais': total_profissionais,
            'total_concluidos': total_concluidos,
            'total_pendentes': total_pendentes,
            'total_em_progresso': total_em_progresso,
            'percentual_conclusao': percentual_conclusao,
            'tempo_medio_conclusao': tempo_medio['tempo_medio'],
            'profissionais_por_status': {
                item['status']: item['total']
                for item in profissionais_por_status
            },
            # Aliases para compatibilidade com templates
            'participantes_concluidos': total_concluidos,
            'participantes_pendentes': total_pendentes,
            'participantes_em_progresso': total_em_progresso,
        }
        
        serializer = ColetaEstatisticasSerializer(data)
        return Response(serializer.data)
    
    @action(detail=True, methods=['get'])
    def exportar_dados_textuais(self, request, pk=None):
        """Exportar dados textuais das respostas concluídas para Excel"""
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        from django.http import HttpResponse
        from datetime import datetime
        
        coleta = self.get_object()
        
        # Buscar respostas concluídas
        respostas = RespostaColeta.objects.filter(
            coleta_profissional__coleta=coleta,
            coleta_profissional__status=StatusPreenchimento.CONCLUIDO
        ).select_related('coleta_profissional__profissional').prefetch_related('respostas_campos__campo')
        
        if not respostas.exists():
            return Response({'error': 'Nenhuma resposta concluída encontrada'}, status=status.HTTP_404_NOT_FOUND)
        
        # Criar workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Dados Textuais"
        
        # Estilo do cabeçalho
        header_fill = PatternFill(start_color="0066CC", end_color="0066CC", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # Preparar cabeçalhos dinâmicos
        headers = ['Profissional', 'CPF']
        
        # Obter todos os campos do modelo (se houver)
        if coleta.modelo:
            campos = list(coleta.modelo.campos.exclude(tipo__in=['foto', 'assinatura']).order_by('ordem'))
            headers.extend([campo.nome_campo for campo in campos])
        
        # Escrever cabeçalho
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        # Dados
        row = 2
        for resposta in respostas:
            nome = resposta.coleta_profissional.profissional.nome if resposta.coleta_profissional and resposta.coleta_profissional.profissional else ''
            cpf = resposta.coleta_profissional.profissional.cpf if resposta.coleta_profissional and resposta.coleta_profissional.profissional else ''

            ws.cell(row=row, column=1, value=nome)
            # CPF: force text formatting and center
            cpf_cell = ws.cell(row=row, column=2, value=cpf)
            cpf_cell.number_format = '@'
            cpf_cell.alignment = Alignment(horizontal='center', vertical='center')

            if coleta.modelo:
                col = 3
                for campo in campos:
                    try:
                        resposta_campo = resposta.respostas_campos.get(campo=campo)
                        valor = resposta_campo.obter_valor()
                        ws.cell(row=row, column=col, value=str(valor) if valor else '')
                    except:
                        ws.cell(row=row, column=col, value='')
                    col += 1

            row += 1

        # Ajustar largura das colunas (aprox. pixels -> excel units)
        from openpyxl.utils import get_column_letter
        widths = {
            1: 57,  # Nome ~400px
            2: 14,  # CPF ~100px
        }
        # Default width for remaining columns
        default_width = 25
        total_cols = len(headers)
        for col in range(1, total_cols + 1):
            w = widths.get(col, default_width)
            ws.column_dimensions[get_column_letter(col)].width = w
        
        # Preparar resposta HTTP
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"dados_textuais_{coleta.nome.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        wb.save(response)
        return response
    
    @action(detail=True, methods=['get'])
    def exportar_participantes_csv(self, request, pk=None):
        """Exportar lista de todos os participantes em CSV"""
        import csv
        from django.http import HttpResponse
        from datetime import datetime
        
        coleta = self.get_object()
        
        response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
        filename = f"participantes_{coleta.nome.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        # Adicionar BOM para Excel reconhecer UTF-8
        response.write('\ufeff')
        
        writer = csv.writer(response)
        writer.writerow(['Nome', 'CPF', 'Email', 'Departamento', 'Status'])
        
        for cp in coleta.coleta_profissionais.select_related('profissional__departamento').all():
            writer.writerow([
                cp.profissional.nome,
                cp.profissional.cpf,
                cp.profissional.email or '',
                cp.profissional.departamento.nome if cp.profissional.departamento else '',
                cp.get_status_display()
            ])
        
        return response

    @action(detail=True, methods=['get'])
    def exportar_participantes_excel(self, request, pk=None):
        """Exportar lista de participantes para Excel com formatação solicitada

        Column width approximations (pixels -> excel width units):
        - Nome: 400px -> ~57
        - CPF: 100px -> ~14
        - Email: 200px -> ~29
        - Status: 100px -> ~14
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment
            from openpyxl.utils import get_column_letter
            from django.http import HttpResponse
            from datetime import datetime

            coleta = self.get_object()

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'Participantes'

            # Header
            header_font = Font(bold=True)
            headers = ['Nome', 'CPF', 'Email', 'Departamento', 'Status']
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')

            # Rows
            row = 2
            for cp in coleta.coleta_profissionais.select_related('profissional__departamento').all():
                nome = cp.profissional.nome if cp.profissional else ''
                cpf = cp.profissional.cpf if cp.profissional else ''
                email = cp.profissional.email or ''
                departamento = cp.profissional.departamento.nome if cp.profissional and cp.profissional.departamento else ''
                status = cp.get_status_display()

                ws.cell(row=row, column=1, value=nome)
                cpf_cell = ws.cell(row=row, column=2, value=cpf)
                # Ensure CPF is treated as text
                cpf_cell.number_format = '@'
                cpf_cell.alignment = Alignment(horizontal='center', vertical='center')

                ws.cell(row=row, column=3, value=email)
                ws.cell(row=row, column=4, value=departamento)
                status_cell = ws.cell(row=row, column=5, value=status)
                status_cell.alignment = Alignment(horizontal='center', vertical='center')

                row += 1

            # Column widths (approximate conversion from pixels to excel width units)
            widths = {
                1: 57,   # Nome ~400px
                2: 14,   # CPF ~100px
                3: 29,   # Email ~200px
                4: 20,   # Departamento (reasonable default)
                5: 14,   # Status ~100px
            }

            for col_idx, width in widths.items():
                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = width

            # Prepare response
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            filename = f"participantes_{coleta.nome.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            wb.save(response)
            return response
        except Exception as e:
            import logging, traceback
            logger = logging.getLogger(__name__)
            logger.exception('Erro ao exportar participantes Excel: %s', e)
            logger.error(traceback.format_exc())
            return Response({'error': f'Erro ao exportar participantes: {str(e)}'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
    
    @action(detail=True, methods=['get'])
    def exportar_midias(self, request, pk=None):
        """Exportar fotos e assinaturas em arquivo ZIP"""
        import zipfile
        import io
        import os
        from django.http import HttpResponse
        from datetime import datetime
        
        coleta = self.get_object()
        
        # Verificar se há foto ou assinatura na coleta
        tem_fotos = RespostaColeta.objects.filter(
            coleta_profissional__coleta=coleta,
            coleta_profissional__status=StatusPreenchimento.CONCLUIDO,
            foto__isnull=False
        ).exists()
        
        tem_assinaturas = RespostaColeta.objects.filter(
            coleta_profissional__coleta=coleta,
            coleta_profissional__status=StatusPreenchimento.CONCLUIDO,
            assinatura__isnull=False
        ).exists()
        
        if not tem_fotos and not tem_assinaturas:
            return Response({'error': 'Nenhuma mídia encontrada'}, status=status.HTTP_404_NOT_FOUND)
        
        # Criar arquivo ZIP em memória
        zip_buffer = io.BytesIO()
        
        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            # Buscar respostas com mídias
            respostas = RespostaColeta.objects.filter(
                coleta_profissional__coleta=coleta,
                coleta_profissional__status=StatusPreenchimento.CONCLUIDO
            ).select_related('coleta_profissional__profissional', 'foto', 'assinatura')
            
            for resposta in respostas:
                cpf = resposta.coleta_profissional.profissional.cpf.replace('.', '').replace('-', '')
                pasta_profissional = f"{cpf}/"
                
                # Adicionar foto se existir
                if hasattr(resposta, 'foto') and resposta.foto.arquivo:
                    try:
                        foto_path = resposta.foto.arquivo.path
                        if os.path.exists(foto_path):
                            extensao = os.path.splitext(foto_path)[1]
                            zip_file.write(foto_path, f"{pasta_profissional}foto{extensao}")
                    except:
                        pass
                
                # Adicionar assinatura se existir
                if hasattr(resposta, 'assinatura') and resposta.assinatura.arquivo:
                    try:
                        assinatura_path = resposta.assinatura.arquivo.path
                        if os.path.exists(assinatura_path):
                            extensao = os.path.splitext(assinatura_path)[1]
                            zip_file.write(assinatura_path, f"{pasta_profissional}assinatura{extensao}")
                    except:
                        pass
        
        # Preparar resposta HTTP
        response = HttpResponse(zip_buffer.getvalue(), content_type='application/zip')
        filename = f"midias_{coleta.nome.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        return response
    
    @action(detail=True, methods=['post'])
    def remover_participante(self, request, pk=None):
        """Remover participante de uma coleta"""
        coleta = self.get_object()
        profissional_id = request.data.get('profissional_id')
        
        if not profissional_id:
            return Response({'error': 'profissional_id é obrigatório'}, status=status.HTTP_400_BAD_REQUEST)
        
        try:
            cp = ColetaProfissional.objects.get(
                coleta=coleta,
                profissional_id=profissional_id
            )
            
            profissional_nome = cp.profissional.nome
            
            # Remover
            cp.delete()
            
            # Atualizar totais
            coleta.atualizar_totais()
            
            # Registrar no log de auditoria
            from auditoria.models import LogAuditoria, TipoAcao
            from empresas.models import Empresa
            
            try:
                empresa = Empresa.objects.get(usuario=request.user)
                LogAuditoria.registrar(
                    usuario=request.user,
                    tipo_acao=TipoAcao.PARTICIPANTE_REMOVIDO,
                    descricao=f"Profissional '{profissional_nome}' removido da coleta '{coleta.nome}'",
                    empresa=empresa,
                    coleta=coleta,
                    dados_adicionais={'profissional_id': str(profissional_id), 'profissional_nome': profissional_nome},
                    request=request
                )
            except Empresa.DoesNotExist:
                pass
            
            return Response({'message': f'Participante {profissional_nome} removido com sucesso.'})
            
        except ColetaProfissional.DoesNotExist:
            return Response({'error': 'Participante não encontrado nesta coleta'}, status=status.HTTP_404_NOT_FOUND)
    
    @action(detail=True, methods=['post'])
    def refazer(self, request, pk=None):
        """Refazer coleta - apagar todas as respostas e reabrir"""
        coleta = self.get_object()
        
        # Verificar se coleta está concluída ou cancelada
        if coleta.status not in ['concluida', 'cancelada']:
            return Response(
                {'error': 'Apenas coletas concluídas ou canceladas podem ser refeitas'}, 
                status=status.HTTP_400_BAD_REQUEST
            )
        
        # Apagar todas as respostas
        from .models import RespostaColeta, Foto, Assinatura
        
        participantes_count = 0
        respostas_count = 0
        
        for cp in coleta.coleta_profissionais.all():
            # Apagar resposta se existir
            if hasattr(cp, 'resposta'):
                resposta = cp.resposta
                
                # Apagar foto
                if hasattr(resposta, 'foto'):
                    resposta.foto.delete()
                
                # Apagar assinatura
                if hasattr(resposta, 'assinatura'):
                    resposta.assinatura.delete()
                
                # Apagar resposta
                resposta.delete()
                respostas_count += 1
            
            # Resetar status do participante
            cp.status = 'pendente'
            cp.data_inicio = None
            cp.data_conclusao = None
            cp.save()
            participantes_count += 1
        
        # Reabrir coleta
        coleta.status = 'em_andamento'
        coleta.data_fim = None
        coleta.save()
        
        # Atualizar totais
        coleta.atualizar_totais()
        
        # Registrar no log
        from auditoria.models import LogAuditoria, TipoAcao
        from empresas.models import Empresa
        
        try:
            empresa = Empresa.objects.get(usuario=request.user)
            LogAuditoria.registrar(
                usuario=request.user,
                tipo_acao=TipoAcao.COLETA_INICIADA,
                descricao=f"Coleta '{coleta.nome}' refeita - {respostas_count} respostas apagadas, {participantes_count} participantes resetados",
                empresa=empresa,
                coleta=coleta,
                dados_adicionais={
                    'respostas_apagadas': respostas_count,
                    'participantes_resetados': participantes_count
                },
                request=request
            )
        except Empresa.DoesNotExist:
            pass
        
        return Response({
            'message': f'Coleta reaberta com sucesso! {respostas_count} respostas apagadas e {participantes_count} participantes podem preencher novamente.'
        })
    
    @action(detail=True, methods=['post'])
    def refazer_individual(self, request, pk=None):
        """Permitir que profissional específico refaça o formulário"""
        coleta = self.get_object()
        profissional_id = request.data.get('profissional_id')
        
        if not profissional_id:
            return Response({'error': 'profissional_id é obrigatório'}, status=status.HTTP_400_BAD_REQUEST)
        
        try:
            cp = ColetaProfissional.objects.get(
                coleta=coleta,
                profissional_id=profissional_id
            )
            
            profissional_nome = cp.profissional.nome
            
            # Apagar resposta se existir
            if hasattr(cp, 'resposta'):
                resposta = cp.resposta
                
                # Apagar foto
                if hasattr(resposta, 'foto'):
                    resposta.foto.delete()
                
                # Apagar assinatura
                if hasattr(resposta, 'assinatura'):
                    resposta.assinatura.delete()
                
                # Apagar resposta
                resposta.delete()
            
            # Resetar status
            cp.status = 'pendente'
            cp.data_inicio = None
            cp.data_conclusao = None
            cp.save()
            
            # Atualizar totais
            coleta.atualizar_totais()
            
            # Se coleta estava concluída, reabrir para em_andamento
            if coleta.status == 'concluida':
                coleta.status = 'em_andamento'
                coleta.data_fim = None
                coleta.save()
            
            # Registrar no log
            from auditoria.models import LogAuditoria, TipoAcao
            from empresas.models import Empresa
            
            try:
                empresa = Empresa.objects.get(usuario=request.user)
                LogAuditoria.registrar(
                    usuario=request.user,
                    tipo_acao=TipoAcao.PARTICIPANTE_ADICIONADO,
                    descricao=f"Profissional '{profissional_nome}' pode refazer formulário da coleta '{coleta.nome}'",
                    empresa=empresa,
                    coleta=coleta,
                    dados_adicionais={'profissional_id': str(profissional_id), 'profissional_nome': profissional_nome},
                    request=request
                )
            except Empresa.DoesNotExist:
                pass
            
            return Response({'message': f'{profissional_nome} pode refazer o formulário!'})
            
        except ColetaProfissional.DoesNotExist:
            return Response({'error': 'Participante não encontrado nesta coleta'}, status=status.HTTP_404_NOT_FOUND)
    
    @action(detail=True, methods=['get'], url_path='participantes/(?P<usuario_id>[^/.]+)/respostas')
    def respostas_participante(self, request, pk=None, usuario_id=None):
        """Retorna respostas de um participante em uma coleta"""
        from empresas.models import Profissional
        try:
            profissional = Profissional.objects.get(usuario=usuario_id)
        except Profissional.DoesNotExist:
            return Response({'detail': 'Participante não encontrado.'}, status=status.HTTP_404_NOT_FOUND)

        # Obter coleta e ColetaProfissional
        coleta = self.get_object()
        try:
            cp = ColetaProfissional.objects.get(coleta=coleta, profissional=profissional)
        except ColetaProfissional.DoesNotExist:
            return Response({'detail': 'Participante não cadastrado nesta coleta.'}, status=status.HTTP_404_NOT_FOUND)

        # Verificar se existe resposta
        if not hasattr(cp, 'resposta'):
            return Response({'detail': 'Nenhuma resposta encontrada para este participante.'}, status=status.HTTP_404_NOT_FOUND)

        resposta = cp.resposta

        # Montar lista de respostas por campo
        itens = []
        for rc in resposta.respostas_campos.select_related('campo').all():
            try:
                valor = rc.obter_valor()
            except Exception:
                valor = None

            itens.append({
                'campo_nome': rc.campo.nome_campo,
                'valor': valor,
                'data_resposta': resposta.created_at
            })

        # Incluir mídia (foto/assinatura) se existir
        if hasattr(resposta, 'foto') and resposta.foto.arquivo:
            itens.append({
                'campo_nome': 'foto',
                'valor': request.build_absolute_uri(resposta.foto.arquivo.url),
                'data_resposta': resposta.created_at
            })
        if hasattr(resposta, 'assinatura') and resposta.assinatura.arquivo:
            itens.append({
                'campo_nome': 'assinatura',
                'valor': request.build_absolute_uri(resposta.assinatura.arquivo.url),
                'data_resposta': resposta.created_at
            })

        return Response(itens, status=status.HTTP_200_OK)


class ColetaProfissionalViewSet(viewsets.ModelViewSet):
    """ViewSet para ColetaProfissional"""
    queryset = ColetaProfissional.objects.select_related(
        'coleta', 'profissional'
    ).all()
    serializer_class = ColetaProfissionalSerializer
    permission_classes = [IsAuthenticated]
    
    def get_queryset(self):
        user = self.request.user
        queryset = ColetaProfissional.objects.select_related(
            'coleta', 'profissional'
        ).all()
        
        # Permissão: profissional só vê coletas da sua empresa
        if user.tipo_usuario == 'profissional':
            from empresas.models import Profissional
            try:
                profissional = Profissional.objects.get(usuario=user)
                queryset = queryset.filter(profissional=profissional, coleta__empresa=profissional.empresa)
            except Profissional.DoesNotExist:
                return ColetaProfissional.objects.none()
        
        # Filtros (support DRF Request or Django request)
        params = getattr(self.request, 'query_params', None)
        if params is None:
            params = getattr(self.request, 'GET', {})

        coleta_id = params.get('coleta_id')
        profissional_id = params.get('profissional_id')
        status_filter = params.get('status')
        
        if coleta_id:
            queryset = queryset.filter(coleta_id=coleta_id)
        if profissional_id:
            queryset = queryset.filter(profissional_id=profissional_id)
        if status_filter:
            queryset = queryset.filter(status=status_filter)
        
        return queryset
    
    @action(detail=True, methods=['post'])
    def iniciar_preenchimento(self, request, pk=None):
        """Iniciar preenchimento"""
        coleta_prof = self.get_object()
        coleta_prof.iniciar()
        return Response({'message': 'Preenchimento iniciado.'})
    
    @action(detail=True, methods=['post'])
    def finalizar_preenchimento(self, request, pk=None):
        """Finalizar preenchimento"""
        coleta_prof = self.get_object()
        coleta_prof.finalizar()
        return Response({'message': 'Preenchimento finalizado.'})

    @action(detail=True, methods=['post'])
    def enviar_notificacao(self, request, pk=None):
        """Endpoint de compatibilidade para enviar notificação (simulado nos testes)."""
        # Implementação mínima para testes - em produção, delegar a Celery
        return Response({'message': 'Notificação enviada (simulada).'})


class RespostaColetaViewSet(viewsets.ModelViewSet):
    """ViewSet para RespostaColeta"""
    queryset = RespostaColeta.objects.select_related('coleta_profissional').all()
    serializer_class = RespostaColetaSerializer
    permission_classes = [IsAuthenticated]
    
    def get_serializer_class(self):
        if self.action == 'create':
            return RespostaColetaCreateSerializer
        return RespostaColetaSerializer
    
    def get_queryset(self):
        queryset = RespostaColeta.objects.select_related('coleta_profissional').all()
        
        # Filtros (support DRF Request or Django request)
        params = getattr(self.request, 'query_params', None)
        if params is None:
            params = getattr(self.request, 'GET', {})

        coleta_profissional_id = params.get('coleta_profissional_id')
        
        if coleta_profissional_id:
            queryset = queryset.filter(coleta_profissional_id=coleta_profissional_id)
        
        return queryset

    @action(detail=True, methods=['post'])
    def validar(self, request, pk=None):
        """Valida uma resposta (compatibilidade com testes)."""
        resposta = self.get_object()
        try:
            valido, mensagem = resposta.validar()
            return Response({'valido': valido, 'mensagem': mensagem})
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


class NotificacaoViewSet(viewsets.ModelViewSet):
    """ViewSet para Notificacao"""
    queryset = Notificacao.objects.select_related('destinatario').all()
    serializer_class = NotificacaoSerializer
    permission_classes = [IsAuthenticated]
    
    def get_queryset(self):
        user = self.request.user
        queryset = Notificacao.objects.select_related('destinatario').all()
        
        # Usuários comuns veem apenas suas notificações
        if user.tipo_usuario not in ['gerente']:
            queryset = queryset.filter(destinatario=user)
        
        # Filtros (support DRF Request or Django request)
        params = getattr(self.request, 'query_params', None)
        if params is None:
            params = getattr(self.request, 'GET', {})

        lido = params.get('lido')
        tipo = params.get('tipo')
        
        if lido is not None:
            queryset = queryset.filter(lido=lido.lower() == 'true')
        if tipo:
            queryset = queryset.filter(tipo=tipo)
        
        return queryset
    
    @action(detail=True, methods=['post'])
    def marcar_lida(self, request, pk=None):
        """Marcar notificação como lida"""
        notificacao = self.get_object()
        notificacao.marcar_lido()
        return Response({'message': 'Notificação marcada como lida.'})
    
    @action(detail=False, methods=['post'])
    def marcar_todas_lidas(self, request):
        """Marcar todas as notificações como lidas"""
        count = Notificacao.objects.filter(
            destinatario=request.user,
            lido=False
        ).update(lido=True)
        return Response({'message': f'{count} notificação(ões) marcada(s) como lida(s).'})

