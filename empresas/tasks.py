import logging
import uuid

import openpyxl
from celery import shared_task
from django.contrib.auth.hashers import make_password
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.utils import timezone

from auditoria.models import LogAuditoria, TipoAcao
from empresas.models import ImportacaoProfissionais, Profissional
from usuarios.models import Usuario


logger = logging.getLogger(__name__)
TAMANHO_LOTE = 500
LIMITE_ERROS_RETORNADOS = 100


def _registrar_erro(erros, linha, erro):
    if len(erros) < LIMITE_ERROS_RETORNADOS:
        erros.append({'linha': linha, 'erro': str(erro)})


@shared_task
def processar_importacao_profissionais(importacao_id):
    """Processa um arquivo Excel já salvo, fora do ciclo de vida da requisição web."""
    importacao = ImportacaoProfissionais.objects.select_related('empresa', 'departamento', 'solicitada_por').get(
        pk=importacao_id
    )
    if importacao.status in {ImportacaoProfissionais.Status.CONCLUIDA, ImportacaoProfissionais.Status.FALHOU}:
        return {'status': importacao.status}
    claimed = ImportacaoProfissionais.objects.filter(
        pk=importacao.pk,
        status=ImportacaoProfissionais.Status.AGUARDANDO,
    ).update(status=ImportacaoProfissionais.Status.PROCESSANDO, iniciada_em=timezone.now())
    if not claimed:
        importacao.refresh_from_db(fields=['status'])
        return {'status': importacao.status}
    importacao.status = ImportacaoProfissionais.Status.PROCESSANDO

    erros = []
    total_erros = 0
    ignorados = 0
    importados = 0
    processados = 0
    try:
        importacao.arquivo.open('rb')
        workbook = openpyxl.load_workbook(importacao.arquivo, read_only=True, data_only=True)
        worksheet = workbook.active
        importacao.total_linhas = max(worksheet.max_row - 1, 0)
        importacao.save(update_fields=['total_linhas'])

        cpfs_existentes = set(Profissional.objects.values_list('cpf', flat=True))
        emails_existentes = set(Usuario.objects.values_list('email', flat=True))
        usuarios = []
        profissionais = []

        def salvar_lote():
            nonlocal importados
            if not usuarios:
                return
            with transaction.atomic():
                Usuario.objects.bulk_create(usuarios, batch_size=TAMANHO_LOTE)
                Profissional.objects.bulk_create(profissionais, batch_size=TAMANHO_LOTE)
            importados += len(profissionais)
            usuarios.clear()
            profissionais.clear()

        for linha, valores in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(valor not in (None, '') for valor in valores):
                continue
            processados += 1
            nome = str(valores[0]).strip() if len(valores) > 0 and valores[0] else ''
            cpf = ''.join(filter(str.isdigit, str(valores[1]))) if len(valores) > 1 and valores[1] else ''
            email = str(valores[2]).strip() if len(valores) > 2 and valores[2] not in (None, '-') else None

            if not nome or len(cpf) != 11:
                _registrar_erro(erros, linha, 'Nome e CPF válido de 11 dígitos são obrigatórios')
                total_erros += 1
                continue
            if cpf in cpfs_existentes:
                ignorados += 1
                continue
            if email:
                try:
                    validate_email(email)
                except ValidationError:
                    _registrar_erro(erros, linha, 'E-mail inválido')
                    total_erros += 1
                    continue

            email_login = email or f'profissional_{cpf}_{uuid.uuid4().hex[:8]}@clickid.temp'
            while email_login in emails_existentes:
                email_login = f'profissional_{cpf}_{uuid.uuid4().hex[:8]}@clickid.temp'

            usuario = Usuario(
                email=email_login,
                tipo_usuario=Usuario.TIPO_PROFISSIONAL,
                password=make_password(cpf[-4:]),
            )
            usuarios.append(usuario)
            profissionais.append(Profissional(
                usuario=usuario,
                empresa=importacao.empresa,
                departamento=importacao.departamento,
                nome=nome,
                cpf=cpf,
                email=email,
                cadastrado_por=importacao.solicitada_por,
            ))
            cpfs_existentes.add(cpf)
            emails_existentes.add(email_login)

            if len(usuarios) >= TAMANHO_LOTE:
                salvar_lote()
                importacao.total_processados = processados
                importacao.total_importados = importados
                importacao.total_ignorados = ignorados
                importacao.total_erros = total_erros
                importacao.erros = erros
                importacao.save(update_fields=['total_processados', 'total_importados', 'total_ignorados', 'total_erros', 'erros'])

        salvar_lote()
        workbook.close()
        importacao.status = ImportacaoProfissionais.Status.CONCLUIDA
        importacao.total_processados = processados
        importacao.total_importados = importados
        importacao.total_ignorados = ignorados
        importacao.total_erros = total_erros
        importacao.erros = erros
        importacao.finalizada_em = timezone.now()
        importacao.save()
        LogAuditoria.registrar(
            usuario=importacao.solicitada_por,
            tipo_acao=TipoAcao.PROFISSIONAL_CRIADO,
            descricao=(f'Importação Excel concluída: {importados} profissionais criados, '
                       f'{ignorados} ignorados, {total_erros} erros'),
            empresa=importacao.empresa,
            dados_adicionais={'importacao_id': str(importacao.pk), 'arquivo': importacao.arquivo.name},
        )
        return {'status': importacao.status, 'total_importados': importados}
    except Exception as exc:
        logger.exception('Falha na importação assíncrona %s', importacao_id)
        importacao.status = ImportacaoProfissionais.Status.FALHOU
        importacao.mensagem_erro = 'Não foi possível processar o arquivo. Tente novamente ou contate o suporte.'
        importacao.finalizada_em = timezone.now()
        importacao.erros = erros
        importacao.total_processados = processados
        importacao.total_importados = importados
        importacao.total_ignorados = ignorados
        importacao.total_erros = total_erros
        importacao.save()
        raise exc
    finally:
        try:
            importacao.arquivo.close()
        except Exception:
            pass
